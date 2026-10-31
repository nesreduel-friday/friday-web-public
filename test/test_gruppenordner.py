"""Gruppenordner über IServ-WebDAV — ohne Netz, ohne Gemini."""

import asyncio
import io
import os
import tempfile
import unittest
from unittest import mock

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from kern import konnektoren
from kern.dienste import webdav
from kern.main import app
from kern.schule import gruppenordner, ki, noten

WAHL = [{"name": "Biologie", "pruefung": 2, "stunden": 5, "halbjahre": ["12/1", "12/2", "13/1", "13/2"]},
        {"name": "Mathematik", "pruefung": 1, "stunden": 5, "halbjahre": ["12/1", "12/2", "13/1", "13/2"]}]


def scan_pdf(farben=("red", "green")) -> bytes:
    """Ein Bild-PDF ohne Textschicht (wie ein eingescanntes Blatt)."""
    seiten = [Image.new("RGB", (300, 400), f) for f in farben]
    puffer = io.BytesIO()
    seiten[0].save(puffer, "PDF", save_all=True, append_images=seiten[1:])
    return puffer.getvalue()


def pptx_bytes(folien: list[str]) -> bytes:
    praesentation = Presentation()
    for text in folien:
        folie = praesentation.slides.add_slide(praesentation.slide_layouts[6])
        box = folie.shapes.add_textbox(0, 0, 1000000, 1000000)
        box.text_frame.text = text
    puffer = io.BytesIO()
    praesentation.save(puffer)
    return puffer.getvalue()


def docx_bytes(absaetze: list[str]) -> bytes:
    dokument = Document()
    for text in absaetze:
        dokument.add_paragraph(text)
    puffer = io.BytesIO()
    dokument.save(puffer)
    return puffer.getvalue()


class Erkennen(unittest.TestCase):
    def test_kurs_von_pfad(self):
        self.assertEqual(gruppenordner.kurs_von_pfad("Gruppenordner/Biologie eA/Blatt3.pdf", "Gruppenordner"),
                         "Biologie eA")
        self.assertEqual(gruppenordner.kurs_von_pfad("Gruppenordner/Mathematik/Unter/Blatt.pdf", "Gruppenordner"),
                         "Mathematik")
        self.assertEqual(gruppenordner.kurs_von_pfad("Gruppenordner/direkt.pdf", "Gruppenordner"), "direkt.pdf")

    def test_typ_von_dateiname(self):
        z = gruppenordner.typ_von_dateiname
        self.assertEqual([z("a.pdf"), z("a.PDF"), z("a.jpg"), z("a.png"), z("a.pptx"), z("a.ppt"),
                          z("a.docx"), z("a.doc"), z("a.zip"), z("ohne_endung")],
                         ["pdf", "pdf", "bild", "bild", "pptx", "pptx", "docx", "docx", "sonstiges", "sonstiges"])

    def test_text_aus_pptx_liest_folien(self):
        text = gruppenordner.text_aus_pptx(pptx_bytes(["Erste Folie", "Zweite Folie"]))
        self.assertIn("Erste Folie", text)
        self.assertIn("Zweite Folie", text)

    def test_text_aus_docx_liest_absaetze(self):
        text = gruppenordner.text_aus_docx(docx_bytes(["Erster Absatz", "Zweiter Absatz"]))
        self.assertIn("Erster Absatz", text)
        self.assertIn("Zweiter Absatz", text)

    def test_scan_pdf_ohne_textschicht_ist_leer(self):
        self.assertEqual(gruppenordner.text_aus_pdf(scan_pdf()), "")

    def test_text_lokal_dispatch(self):
        with mock.patch.object(gruppenordner, "text_aus_pdf", return_value="genug Text " * 10), \
             mock.patch.object(gruppenordner, "text_aus_pptx", return_value="pptx-text"), \
             mock.patch.object(gruppenordner, "text_aus_docx", return_value="docx-text"):
            self.assertEqual(gruppenordner.text_lokal("pdf", b""), "genug Text " * 10)
            self.assertEqual(gruppenordner.text_lokal("pptx", b""), "pptx-text")
            self.assertEqual(gruppenordner.text_lokal("docx", b""), "docx-text")

    def test_text_lokal_kurzer_pdf_text_gilt_als_scan(self):
        with mock.patch.object(gruppenordner, "text_aus_pdf", return_value="kurz"):
            self.assertEqual(gruppenordner.text_lokal("pdf", b""), "")

    def test_text_lokal_bild_liefert_nichts(self):
        self.assertEqual(gruppenordner.text_lokal("bild", b""), "")

    def test_text_lokal_kaputte_datei_bricht_nicht(self):
        self.assertEqual(gruppenordner.text_lokal("pptx", b"kein echtes pptx"), "")


class Kurse(unittest.TestCase):
    FAECHER = [{"id": i, "name": n} for i, n in enumerate(
        ["Biologie", "Chemie", "Mathematik", "Deutsch", "Englisch", "Philosophie", "Politik", "Sport", "Seminarfach"], 1)]

    def test_kurscode(self):
        self.assertEqual(gruppenordner.kurs_code("ABI28 Kurs 12 EA BI281"), "BI281")
        self.assertEqual(gruppenordner.kurs_code("ABI28 SP 12  SP282"), "SP282")
        for kein in ("Schüler", "Benutzer", "Schüler Jahrgang 12", "Projekt Reporter"):
            self.assertEqual(gruppenordner.kurs_code(kein), "", kein)

    def test_fach_ueber_kurscode(self):
        z = lambda k: {f["id"]: f["name"] for f in self.FAECHER}.get(gruppenordner.kurs_fach_id(k, self.FAECHER))
        self.assertEqual([z("ABI28 Kurs 12 EA BI281"), z("ABI28 Kurs 12 EA CH281"), z("ABI28 Kurs 12 EA MA281"),
                          z("ABI28 Kurs 12 GA PL281"), z("ABI28 Kurs 12 GA PO281"), z("ABI28 SP 12  SP282"),
                          z("ABI28 Seminarfach SF283")],
                         ["Biologie", "Chemie", "Mathematik", "Philosophie", "Politik", "Sport", "Seminarfach"])

    def test_standardwurzel(self):
        with mock.patch.dict(os.environ, {"GRUPPENORDNER_ORDNER": ""}):
            self.assertEqual(gruppenordner.wurzel(), "Groups")


class Mit_Ordner(unittest.TestCase):
    def setUp(self):
        self.env = mock.patch.dict(os.environ, {"FRIDAY_DATEN": tempfile.mkdtemp(),
                                                "ISERV_USER": "eike", "ISERV_PASS": "geheim"})
        self.env.start()
        self.addCleanup(self.env.stop)
        webdav.vergessen()
        self.addCleanup(webdav.vergessen)
        gruppenordner.zustand.update(phase="", fehler=None, ergebnis=None, letzter=None)
        r = TestClient(app).put("/api/noten/fachwahl", json={"faecher": WAHL})
        self.assertEqual(r.status_code, 200, r.text)
        self.faecher = noten.standard().uebersicht()["faecher"]

    def einschalten(self):
        konnektoren.setzen("schulmanager", True)
        konnektoren.setzen("iserv_gruppenordner", True)


class Abgleichen(Mit_Ordner):
    def test_neue_dateien_werden_eingetragen_und_dem_fach_zugeordnet(self):
        a = gruppenordner.standard()
        dateien = [{"pfad": "Gruppenordner/Biologie eA/Blatt.pdf", "ordner": False, "groesse": 100,
                    "geaendert": "2026-09-01T00:00:00+00:00"}]
        zu_laden = a.abgleichen(dateien, "Gruppenordner", self.faecher)
        self.assertEqual(len(zu_laden), 1)
        bio_id = next(f["id"] for f in self.faecher if f["name"] == "Biologie")
        self.assertEqual((zu_laden[0]["kurs_text"], zu_laden[0]["fach_id"], zu_laden[0]["dateiname"], zu_laden[0]["typ"]),
                         ("Biologie eA", bio_id, "Blatt.pdf", "pdf"))

    def test_verschwundene_datei_wird_als_weg_markiert(self):
        a = gruppenordner.standard()
        d = [{"pfad": "Gruppenordner/X/a.pdf", "ordner": False, "groesse": 1, "geaendert": "2026-09-01T00:00:00+00:00"}]
        a.abgleichen(d, "Gruppenordner", self.faecher)
        a.abgleichen([], "Gruppenordner", self.faecher)
        self.assertEqual(a.alle(), [])
        self.assertEqual(len(a.alle(mit_weg=True)), 1)

    def test_bereits_gelesene_unveraenderte_datei_wird_nicht_erneut_verlangt(self):
        a = gruppenordner.standard()
        d = [{"pfad": "Gruppenordner/X/a.pdf", "ordner": False, "groesse": 100, "geaendert": "2026-09-01T00:00:00+00:00"}]
        [e] = a.abgleichen(d, "Gruppenordner", self.faecher)
        a.setzen(e["id"], status="gelesen", gelesen_groesse=100, gelesen_geaendert="2026-09-01T00:00:00+00:00")
        self.assertEqual(a.abgleichen(d, "Gruppenordner", self.faecher), [])

    def test_geaenderte_datei_wird_erneut_verlangt(self):
        a = gruppenordner.standard()
        d = [{"pfad": "Gruppenordner/X/a.pdf", "ordner": False, "groesse": 100, "geaendert": "2026-09-01T00:00:00+00:00"}]
        [e] = a.abgleichen(d, "Gruppenordner", self.faecher)
        a.setzen(e["id"], status="gelesen", gelesen_groesse=100, gelesen_geaendert="2026-09-01T00:00:00+00:00")
        d2 = [{"pfad": "Gruppenordner/X/a.pdf", "ordner": False, "groesse": 200, "geaendert": "2026-09-02T00:00:00+00:00"}]
        wieder = a.abgleichen(d2, "Gruppenordner", self.faecher)
        self.assertEqual(len(wieder), 1)

    def test_nicht_unterstuetzte_datei_wird_nicht_dauernd_erneut_verlangt(self):
        a = gruppenordner.standard()
        d = [{"pfad": "Gruppenordner/X/a.zip", "ordner": False, "groesse": 100, "geaendert": "2026-09-01T00:00:00+00:00"}]
        [e] = a.abgleichen(d, "Gruppenordner", self.faecher)
        a.setzen(e["id"], status="nicht_unterstuetzt", gelesen_groesse=100, gelesen_geaendert="2026-09-01T00:00:00+00:00")
        self.assertEqual(a.abgleichen(d, "Gruppenordner", self.faecher), [])


class MitLauf(Mit_Ordner):
    def lauf(self, dateien: dict[str, bytes], gemini):
        """dateien: pfad → Rohbytes."""
        liste = [{"pfad": p, "ordner": False, "groesse": len(b), "geaendert": "2026-09-01T00:00:00+00:00"}
                 for p, b in dateien.items()]

        async def laden(pfad, ziel):
            ziel.parent.mkdir(parents=True, exist_ok=True)
            ziel.write_bytes(dateien[pfad])
            return len(dateien[pfad])

        kurse = sorted({"/".join(p.split("/")[:2]) for p in dateien} | {"Groups/Schüler"})   # „Schüler" ist kein Kurs

        async def auflisten(kurs):
            return [d for d in liste if d["pfad"].startswith(kurs + "/")]

        with mock.patch.object(webdav, "unterordner", mock.AsyncMock(return_value=kurse)), \
             mock.patch.object(webdav, "auflisten", mock.AsyncMock(side_effect=auflisten)), \
             mock.patch.object(webdav, "laden", mock.AsyncMock(side_effect=laden)), \
             mock.patch.object(ki, "json_anfrage", gemini):
            return asyncio.run(gruppenordner.durchgang(pause=0))


class Durchgang(MitLauf):
    def test_pptx_wird_lokal_gelesen_und_nur_zusammengefasst(self):
        gemini = mock.AsyncMock(return_value=({"zusammenfassung": "Kurz.", "themen": ["a", "b"],
                                                "wichtiges": ["Q = m · c · ΔT"]}, "probe"))
        self.lauf({"Groups/ABI28 Kurs 12 EA MA281/Folien.pptx": pptx_bytes(["Ableitungen", "Kettenregel"])}, gemini)
        [e] = gruppenordner.standard().alle()
        self.assertEqual(e["status"], "gelesen")
        self.assertIn("Ableitungen", e["text"])
        self.assertEqual((e["zusammenfassung"], e["themen"], e["wichtiges"]),
                         ("Kurz.", ["a", "b"], ["Q = m · c · ΔT"]))
        # Kein Bild ging mit — die Anfrage war reiner Text.
        teile = gemini.call_args.args[0]
        self.assertTrue(all("inlineData" not in t for t in teile))

    def test_wichtiges_fehlt_in_der_antwort_wird_leere_liste(self):
        gemini = mock.AsyncMock(return_value=({"zusammenfassung": "Kurz.", "themen": []}, "probe"))
        self.lauf({"Groups/ABI28 Kurs 12 EA MA281/Folien.pptx": pptx_bytes(["x"])}, gemini)
        [e] = gruppenordner.standard().alle()
        self.assertEqual(e["wichtiges"], [])

    def test_scan_pdf_geht_als_bild_an_gemini(self):
        gemini = mock.AsyncMock(return_value=({"text": "Abschrift", "zusammenfassung": "Kurz.", "themen": ["x"]}, "probe"))
        self.lauf({"Groups/ABI28 Kurs 12 EA BI281/Scan.pdf": scan_pdf()}, gemini)
        [e] = gruppenordner.standard().alle()
        self.assertEqual((e["status"], e["text"]), ("gelesen", "Abschrift"))
        teile = gemini.call_args.args[0]
        self.assertTrue(any("inlineData" in t for t in teile))

    def test_bild_datei_geht_an_gemini(self):
        bild = io.BytesIO()
        Image.new("RGB", (100, 100), "red").save(bild, "JPEG")
        gemini = mock.AsyncMock(return_value=({"text": "Tafelbild", "zusammenfassung": "Kurz.", "themen": []}, "probe"))
        self.lauf({"Groups/ABI28 Kurs 12 EA CH281/Foto.jpg": bild.getvalue()}, gemini)
        [e] = gruppenordner.standard().alle()
        self.assertEqual((e["status"], e["text"]), ("gelesen", "Tafelbild"))

    def test_nicht_unterstuetzte_datei_wird_nur_gelistet(self):
        gemini = mock.AsyncMock()
        self.lauf({"Groups/ABI28 Kurs 12 EA BI281/Video.mp4": b"nichtssagend"}, gemini)
        [e] = gruppenordner.standard().alle()
        self.assertEqual(e["status"], "nicht_unterstuetzt")
        self.assertFalse(gruppenordner.standard().original_pfad(e["id"], "mp4").exists())   # gar nicht erst geladen
        gemini.assert_not_awaited()

    def test_kontingent_bricht_den_durchgang_ab_naechster_macht_weiter(self):
        gemini = mock.AsyncMock(side_effect=ki.KiFehler("Gemini ist gerade überlastet"))
        ergebnis = self.lauf({"Groups/ABI28 Kurs 12 EA BI281/a.pptx": pptx_bytes(["x"])}, gemini)
        self.assertEqual(ergebnis["gelesen"], 0)
        self.assertIn("überlastet", gruppenordner.zustand["fehler"])
        [e] = gruppenordner.standard().alle()
        self.assertEqual(e["status"], "fehler")   # nächster Durchgang versucht es wieder (offene() nimmt 'fehler')

    def test_ordner_ohne_kurscode_werden_nicht_aufgelistet(self):
        gemini = mock.AsyncMock()
        with mock.patch.object(webdav, "unterordner", mock.AsyncMock(return_value=["Groups/Schüler", "Groups/Benutzer"])), \
             mock.patch.object(webdav, "auflisten", mock.AsyncMock()) as auflisten:
            asyncio.run(gruppenordner.durchgang(pause=0))
        auflisten.assert_not_awaited()

    def test_original_datei_route(self):
        gemini = mock.AsyncMock(return_value=({"zusammenfassung": "Kurz.", "themen": []}, "probe"))
        self.lauf({"Groups/ABI28 Kurs 12 EA MA281/Folien.pptx": pptx_bytes(["x"])}, gemini)
        [e] = gruppenordner.standard().alle()
        r = TestClient(app).get(f"/api/gruppenordner/{e['id']}/original")
        self.assertEqual(r.status_code, 200)
        self.assertIn("Folien.pptx", r.headers["content-disposition"])

    def test_suche_findet_ueber_zusammenfassung(self):
        gemini = mock.AsyncMock(return_value=({"zusammenfassung": "Ableitungsregeln erklärt.", "themen": ["Analysis"]}, "probe"))
        self.lauf({"Groups/ABI28 Kurs 12 EA MA281/Folien.pptx": pptx_bytes(["Ableitungen"])}, gemini)
        treffer = gruppenordner.standard().suchen("Ableitungsregeln")
        self.assertEqual(len(treffer), 1)
        self.assertEqual(treffer[0]["dateiname"], "Folien.pptx")


class Route(Mit_Ordner):
    def test_route_ohne_freigabe_lehnt_ab(self):
        c = TestClient(app)
        self.assertEqual(c.post("/api/gruppenordner/abgleichen").status_code, 403)
        self.assertEqual(c.get("/api/gruppenordner/suche?q=x").status_code, 403)

    def test_original_datei_route_unbekannt_gibt_404(self):
        self.assertEqual(TestClient(app).get("/api/gruppenordner/999999/original").status_code, 404)

    def test_route_zeigt_kurse_gruppiert(self):
        self.einschalten()
        a = gruppenordner.standard()
        d = [{"pfad": "Gruppenordner/Biologie eA/a.pdf", "ordner": False, "groesse": 1,
              "geaendert": "2026-09-01T00:00:00+00:00"}]
        a.abgleichen(d, "Gruppenordner", self.faecher)
        d = TestClient(app).get("/api/gruppenordner").json()
        self.assertEqual(d["frei"], True)
        self.assertEqual(len(d["kurse"]), 1)
        self.assertEqual((d["kurse"][0]["kurs"], d["kurse"][0]["fach"], d["kurse"][0]["dateien"]),
                         ("Biologie eA", "Biologie", 1))


if __name__ == "__main__":
    unittest.main()
